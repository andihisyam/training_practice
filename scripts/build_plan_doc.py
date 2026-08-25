from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUTPUT_DIR = ROOT / "outputs" / "doc_plan_qa"
DOCX_PATH = DOCS_DIR / "Struktur_Plan_Penelitian_DM_XAI.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))


def style_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, after=6, before=0, line=1.2):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line


def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, after=4, line=1.2)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        style_run(first, bold=True)
        rest = p.add_run(text[len(bold_prefix) :])
        style_run(rest)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=3, line=1.15)
    r = p.add_run(text)
    style_run(r, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=14, after=6, line=1.0)
        r = p.add_run(text)
        style_run(r, size=15, bold=True, color="2E74B5")
    else:
        style_paragraph(p, before=10, after=4, line=1.0)
        r = p.add_run(text)
        style_run(r, size=12, bold=True, color="1F4D78")
    return p


def add_table(doc):
    rows = [
        ("1. Audit & Integrasi", "Cek struktur data, join semua tabel ke patient, validasi label dan PatientGuid.", "Dataset gabungan yang konsisten."),
        ("2. Prepare Data", "Tangani zero semu, missing values, encoding, scaling, dan pembagian train-test.", "Data siap eksperimen."),
        ("3. EDA Terarah", "Lihat imbalance, distribusi fitur, korelasi, sparsity medication, dan perbedaan DM vs non-DM.", "Hipotesis fitur dan risiko data."),
        ("4. Feature Set", "Buat skenario core, clinical, dan full features.", "Beberapa versi dataset untuk diuji."),
        ("5. Baseline Model", "Latih Logistic Regression, SVM, dan KNN sebagai pembanding awal.", "Skor dasar untuk perbandingan."),
        ("6. Boosting Model", "Latih Gradient Boosting, XGBoost, dan LightGBM.", "Kandidat model utama."),
        ("7. Tuning", "Optimasi hyperparameter dengan Optuna pada data train menggunakan cross-validation.", "Model versi terbaik per algoritma."),
        ("8. Evaluasi", "Ukur confusion matrix, accuracy, precision, recall, F1, ROC-AUC, dan PR-AUC.", "Perbandingan performa final."),
        ("9. XAI", "Jalankan SHAP pada model terbaik untuk global dan local explanation.", "Interpretasi faktor penting."),
        ("10. Kesimpulan", "Bandingkan hasil, cek leakage, dan rumuskan model terbaik serta keterbatasannya.", "Narasi hasil penelitian."),
    ]

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.45, 3.45, 1.60]

    hdr = table.rows[0].cells
    headers = ["Tahap", "Fokus Kerja", "Output"]
    for idx, text in enumerate(headers):
        set_cell_width(hdr[idx], widths[idx])
        set_cell_shading(hdr[idx], "E8EEF5")
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(text)
        style_run(r, size=10.5, bold=True, color="1F3A5F")

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_width(cells[idx], widths[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.1)
            r = p.add_run(text)
            style_run(r, size=10)


def build_doc():
    DOCS_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(title, after=3, line=1.0)
    r = title.add_run("Struktur Plan Penelitian")
    style_run(r, size=22, bold=False)

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=10, line=1.1)
    r = subtitle.add_run("Prediksi Diabetes Mellitus Tipe 2 dengan Machine Learning dan Explainable AI")
    style_run(r, size=11, color="555555")

    add_body_paragraph(
        doc,
        "Dokumen ini merangkum alur kerja dari tahap prepare data sampai modeling dan interpretasi XAI, serta daftar uji coba yang perlu dilakukan agar eksperimen rapi, terukur, dan mudah dibandingkan.",
    )

    add_heading(doc, "Alur Kerja Utama", level=1)
    add_table(doc)

    add_heading(doc, "Uji Coba yang Wajib Dilakukan", level=1)
    add_heading(doc, "A. Uji pada Data Preparation", level=2)
    for text in [
        "Bandingkan data sebelum dan sesudah penanganan zero semu seperti BMI_Min = 0 dan Weight_Min = 0.",
        "Uji strategi imputasi yang sederhana dan konsisten, terutama untuk fitur numerik dan fitur indikator.",
        "Bandingkan skenario dengan medication dan tanpa medication karena tabel ini tidak lengkap untuk semua pasien.",
        "Pastikan scaling hanya dipakai untuk model yang membutuhkannya seperti Logistic Regression, SVM, dan KNN.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "B. Uji pada Feature Set", level=2)
    for text in [
        "Core features: patient + transcript.",
        "Clinical features: core + diagnosis + physician specialty.",
        "Full features: semua fitur termasuk medication.",
        "Opsional: buang fitur yang terlalu jarang muncul atau near-zero variance.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "C. Uji pada Model", level=2)
    for text in [
        "Baseline: Logistic Regression, SVM, KNN.",
        "Boosting: Gradient Boosting, XGBoost, LightGBM.",
        "Bandingkan parameter default vs hasil tuning Optuna.",
        "Uji class imbalance handling seperti class_weight atau SMOTE di data train saja.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "D. Uji pada Evaluasi", level=2)
    for text in [
        "Gunakan stratified train-test split dan cross-validation pada data train.",
        "Bandingkan accuracy, precision, recall, F1-score, ROC-AUC, dan PR-AUC.",
        "Fokus utama pemilihan model adalah recall pasien diabetes, lalu F1-score dan PR-AUC.",
        "Uji threshold default 0.5 vs threshold yang dioptimalkan untuk recall atau F1.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "E. Uji pada XAI", level=2)
    for text in [
        "Gunakan SHAP pada model terbaik untuk global importance dan local explanation.",
        "Buat summary plot untuk melihat fitur paling dominan.",
        "Buat dependence plot untuk fitur klinis utama seperti BMI, tekanan darah, atau diagnosis tertentu.",
        "Ambil beberapa contoh pasien untuk force plot agar prediksi per individu bisa dijelaskan.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Catatan Penting", level=1)
    for text in [
        "Waspadai leakage bila label diabetes sangat dekat dengan fitur diagnosis atau medication.",
        "Simpan semua hasil eksperimen dalam tabel perbandingan yang seragam agar keputusan model tidak bias.",
        "Output akhir minimal terdiri dari dataset final, tabel eksperimen, model terbaik, dan interpretasi SHAP.",
    ]:
        add_bullet(doc, text)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(p, after=0, line=1.0)
    r = p.add_run("PracticeFusion Research Plan")
    style_run(r, size=9, color="666666")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
